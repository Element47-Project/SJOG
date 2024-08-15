import os
import io
import pandas as pd
import pyodbc
from dotenv import load_dotenv
import openpyxl
import xlrd
from Gas_csv_Formatting import consumption
from Elec_csv_Formatting import e_formatting
from Gas_billing_csv_Formatting import gb_formatting
import requests
from datetime import datetime, timedelta
from sqlalchemy import create_engine
import tabula
from Apollo import upload_apollo

# Load environment variables
load_dotenv()

# All settings in the .env file, including SQL information.
SQL_SERVER = os.environ.get('AZURE_SQL_SERVER')
SQL_DB_NAME = os.environ.get('AZURE_SQL_DB_NAME')
SQL_USERNAME = os.environ.get('AZURE_SQL_USERNAME')
SQL_PASSWORD = os.environ.get('AZURE_SQL_PASSWORD')
FILE_DIR = os.environ.get('FILE_ADDRESS')

print(f"File directory: {FILE_DIR}")

engine = create_engine(f'mssql+pyodbc://{SQL_USERNAME}:{SQL_PASSWORD}@{SQL_SERVER}/{SQL_DB_NAME}'
                       f'?driver=ODBC+Driver+18+for+SQL+Server')
CONNECTION_STRING = (
    f'DRIVER={{ODBC Driver 18 for SQL Server}};SERVER={SQL_SERVER};'
    f'DATABASE={SQL_DB_NAME};UID={SQL_USERNAME};PWD={SQL_PASSWORD}'
)


def connect_to_db(conn_str):
    """Establishes a connection to the database."""
    with pyodbc.connect(conn_str) as conn:
        cursor = conn.cursor()
    return conn, cursor


def get_all_table_primary_keys(cursor):
    """
    Retrieves all primary key columns for each table in the database.
    """
    # Retrieve all base tables
    cursor.execute("SELECT TABLE_NAME FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_TYPE = 'BASE TABLE'")
    table_names = [row.TABLE_NAME for row in cursor.fetchall()]
    # Dictionary to hold table primary key information
    primary_keys = {}

    for table_name in table_names:
        # Query to find primary key columns for the current table
        pk_query = """
            SELECT COLUMN_NAME
            FROM INFORMATION_SCHEMA.KEY_COLUMN_USAGE
            WHERE TABLE_NAME = ? AND 
            OBJECTPROPERTY(OBJECT_ID(CONSTRAINT_SCHEMA + '.' + CONSTRAINT_NAME), 'IsPrimaryKey') = 1
            ORDER BY ORDINAL_POSITION
        """
        cursor.execute(pk_query, [table_name])
        pk_columns = cursor.fetchall()
        # Add table primary key info to the dictionary
        primary_keys[table_name] = [col.COLUMN_NAME for col in pk_columns]
    return primary_keys, table_names


def get_all_table_columns(cursor):
    """Get all table columns name from Azure SQL."""
    tables_columns = {}
    cursor.execute("SELECT TABLE_NAME, COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS ORDER BY TABLE_NAME, "
                   "ORDINAL_POSITION")
    for row in cursor.fetchall():
        table_name, column_name = row
        if table_name not in tables_columns:
            tables_columns[table_name] = []
        tables_columns[table_name].append(column_name)
    return tables_columns


def fetch_latest_date_from_azure(cursor, table_dict, table_name='Temperature_hourly'):
    primary_key = table_dict[table_name]
    pk_col = primary_key[0]
    cursor.execute(f"SELECT TOP 1 [{pk_col}] FROM [{table_name}] ORDER BY [{pk_col}] DESC")
    last_record = cursor.fetchone()
    return last_record[0] if last_record else None


def process_xlsx_file(file_path, table_dict, cursor):
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True)

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            skip_rows = 0
            for row in sheet.iter_rows(min_row=1, max_col=1):
                cell = row[0]
                if cell.value is not None:
                    skip_rows = cell.row - 1
                    break

            batch_df = pd.read_excel(file_path, sheet_name=sheet_name, skiprows=skip_rows)

            # Find the matching table name based on the column names
            for table_name, azure_columns in get_all_table_columns(cursor).items():
                if all(col in batch_df.columns for col in azure_columns):
                    upload_dataframe_to_azure_sql(batch_df, table_name, cursor, table_dict)
                    delete_file(file_path)
                    break
            else:
                print(f"No matching table found for sheet: {sheet_name}")

    except xlrd.biffh.XLRDError as e:
        if str(e) == "Workbook is encrypted":
            print(f"Cannot process encrypted file: {file_path}")
        else:
            raise


def process_csv_file(file_path, table_dict, cursor):
    file_name_without_extension = os.path.basename(file_path).rsplit('.', 1)[0]
    csv_header = pd.read_csv(file_path, nrows=0).columns.tolist()
    csv_data = pd.read_csv(file_path)
    if file_name_without_extension in table_dict:
        upload_dataframe_to_azure_sql(csv_data, file_name_without_extension, cursor, table_dict)
        delete_file(file_path)
    elif 'CONSUMPTION_HR01' in csv_header:
        df_csv = consumption(csv_data)
        upload_dataframe_to_azure_sql(df_csv, 'TestingGas', cursor, table_dict)
        delete_file(file_path)
    elif 'GAS (GJ)' in csv_header:
        upload_dataframe_to_azure_sql(csv_data, 'TestingGas', cursor, table_dict)
        delete_file(file_path)
    elif 'Unit Of Measure' in csv_header:
        df_csv = e_formatting(csv_data)
        upload_dataframe_to_azure_sql(df_csv, 'TestingElecBilling', cursor, table_dict)
        delete_file(file_path)
    elif 'NET_CHARGE' in csv_header:
        df_csv = gb_formatting(csv_data)
        upload_dataframe_to_azure_sql(df_csv, 'TestingGasBill', cursor, table_dict)
        delete_file(file_path)
    elif 'LogRecNum' in csv_header:
        try:
            df_csv = upload_apollo(csv_data, file_name_without_extension)
            df_csv.to_sql('Apollo_5MINS', engine, if_exists='append', index=False)
            print("Insert Successful")
            delete_file(file_path)
        except pyodbc.Error as e:
            print(e)
    else:
        print("The CSV file cannot be inserted into the Azure SQL DB")


def process_pdf_file(file_path, table_dict, cursor):
    try:
        tables = tabula.read_pdf(file_path, pages='all', multiple_tables=True)
        correct_headers = ['Water Use Year', 'Read Date', 'Reading', 'Dial Reading', 'Kilolitres Used',
                           'Consumption Year to Date', 'Daily Rate']
        for table in tables:
            if not table.empty:
                table.columns = correct_headers
                processed_table = table.iloc[2:].reset_index(drop=True)
                processed_table['Read Date'] = pd.to_datetime(processed_table['Read Date'], format='%d/%m/%Y')
                upload_dataframe_to_azure_sql(processed_table, 'TestingWater', cursor, table_dict)
                delete_file(file_path)

    except Exception as e:
        print(f"Error processing PDF file: {file_path}")
        print(e)


def process_files_in_directory(directory, cursor, table_dict):
    for filename in os.listdir(directory):
        file_path = os.path.join(directory, filename)
        _, extension = os.path.splitext(filename)
        print(f"Processing file: {filename}")
        if extension in ['.xlsx', '.xls', '.xlsm']:
            process_xlsx_file(file_path, table_dict, cursor)
        elif extension == '.csv':
            process_csv_file(file_path, table_dict, cursor)
        elif extension == '.pdf':
            process_pdf_file(file_path, table_dict, cursor)


def upload_dataframe_to_azure_sql(df, table_name, cursor, table_dict):
    print(f"Uploading Data to {table_name}. Please Wait...")
    primary_keys = table_dict[table_name]

    if table_name == 'TestingBilling':
        try:
            df.to_sql(table_name, engine, if_exists='append', index=False)
            print("Insert Successful")
        except pyodbc.Error as e:
            print(e)
        return 1

    if len(primary_keys) == 1:
        pk_col = primary_keys[0]
        cursor.execute(f"SELECT TOP 1 [{pk_col}] FROM [{table_name}] ORDER BY [{pk_col}] DESC")
        last_record = cursor.fetchone()
        last_value = pd.to_datetime(last_record[0], errors='coerce') if last_record else None
        df[pk_col] = pd.to_datetime(df[pk_col], format='%d-%b-%Y %H:%M:%S', errors='coerce')
        if last_value is not None:
            df = df[df[pk_col] > last_value].copy()

    elif len(primary_keys) == 2 and 'END INTERVAL' in primary_keys:
        cursor.execute(f"""
                SELECT [NMI], MAX([END INTERVAL]) 
                FROM [{table_name}] 
                GROUP BY [NMI]
                """)
        last_records = {nmi: max_end_interval for nmi, max_end_interval in cursor.fetchall()}
        filtered_df = pd.DataFrame()
        df['END INTERVAL'] = pd.to_datetime(df['END INTERVAL'], errors='coerce')
        nmi_not_in_last_records = ~df['NMI'].astype(str).isin(last_records.keys())
        temp_df1_indices = df[nmi_not_in_last_records].index
        temp_df1 = df.loc[temp_df1_indices]
        for nmi, last_time in last_records.items():
            temp_indices = df[(df['NMI'].astype(str) == str(nmi)) & (df['END INTERVAL'] > last_time)].index
            temp_df = df.loc[temp_indices]
            filtered_df = pd.concat([filtered_df, temp_df], ignore_index=True)
        df = pd.concat([temp_df1, filtered_df], ignore_index=True)

    elif len(primary_keys) == 2 and 'BILLING PERIOD START DATE' in primary_keys:
        cursor.execute(f"""
                SELECT [NMI], MAX([BILLING PERIOD START DATE]) 
                FROM [{table_name}] 
                GROUP BY [NMI]
                """)
        last_records = {nmi: max_end_interval for nmi, max_end_interval in cursor.fetchall()}
        filtered_df = pd.DataFrame()
        df['BILLING PERIOD START DATE'] = pd.to_datetime(df['BILLING PERIOD START DATE'], errors='coerce')
        nmi_not_in_last_records = ~df['NMI'].astype(str).isin(last_records.keys())
        temp_df1_indices = df[nmi_not_in_last_records].index
        temp_df1 = df.loc[temp_df1_indices]
        for nmi, last_time in last_records.items():
            last_time_datetime = pd.to_datetime(last_time)
            temp_indices = df[(df['NMI'].astype(str) == str(nmi)) & (
                        df['BILLING PERIOD START DATE'] > last_time_datetime)].index
            temp_df = df.loc[temp_indices]
            filtered_df = pd.concat([filtered_df, temp_df], ignore_index=True)
        df = pd.concat([temp_df1, filtered_df], ignore_index=True)

    if df.empty:
        print("No New Rows to Insert After Filtering with Last Records.")
        return

    try:
        df.to_sql(table_name, engine, if_exists='append', index=False)
        print("Insert Successful")
    except pyodbc.Error as e:
        print(e)


def delete_file(file_path):
    """Delete the file at the specified file path."""
    try:
        os.remove(file_path)
        print(f"Deleted file: {file_path}")
    except OSError as e:
        print(f"Error deleting file: {file_path}")
        print(e)


def fetch_weather_data(latitude, longitude, start_date, end_date):
    url = "https://archive-api.open-meteo.com/v1/archive"
    params = {
        "latitude": latitude,
        "longitude": longitude,
        "start_date": start_date,
        "end_date": end_date,
        "hourly": "temperature_2m",
        "timezone": "auto"
    }
    response = requests.get(url, params=params)
    if response.status_code == 200:
        return response.json()
    else:
        return None


def process_weather_data(weather_data):
    # Accessing nested data within 'hourly' key
    hourly_data = weather_data['hourly']
    hourly_times = hourly_data['time']
    hourly_temperatures = hourly_data['temperature_2m']
    # Convert the ISO8601 time strings to datetime objects
    times = pd.to_datetime(hourly_times)
    # Create the DataFrame
    df_weather = pd.DataFrame({'Date_Time': times, 'Temperature': hourly_temperatures})
    return df_weather


def main():
    print("Connecting to SQL Database...")
    # Connect to the Azure SQL
    conn, cursor = connect_to_db(CONNECTION_STRING)
    print("Connected. Loading the Information from Database...")
    table_dict, all_tables = get_all_table_primary_keys(cursor)

    # Process files in the specified directory
    process_files_in_directory(FILE_DIR, cursor, table_dict)

    # Upload the Temperature data
    print("Uploading the Recent Temperature to Azure. Please Wait..")
    table_name = 'Temperature_hourly'
    latest_date = fetch_latest_date_from_azure(cursor, table_dict, table_name)
    if latest_date:
        start_date = (latest_date - timedelta(days=1 / 3)).strftime("%Y-%m-%d")
    else:
        start_date = "2020-01-01"
    end_date = (datetime.now() - timedelta(days=3)).strftime("%Y-%m-%d")
    latitude = -31.9522
    longitude = 115.8614
    weather_data = fetch_weather_data(latitude, longitude, start_date, end_date)
    df_weather = process_weather_data(weather_data)
    upload_dataframe_to_azure_sql(df_weather, table_name, cursor, table_dict)


# Entry
if __name__ == "__main__":
    main()
